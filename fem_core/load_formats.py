from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from fem_core.errors import FemCoreError
from fem_core.text import decode_engineering_text

MAX_LOAD_BYTES = 20 * 1024 * 1024
MAX_ROWS = 200_000
MAX_XLSX_EXTRACTED_BYTES = 100 * 1024 * 1024
MAX_XLSX_MEMBERS = 10_000
TABLE_SUFFIXES = frozenset({".csv", ".txt", ".dat", ".xlsx"})
PEER_SUFFIXES = frozenset({".at1", ".at2"})
SUPPORTED_LOAD_SUFFIXES = TABLE_SUFFIXES | PEER_SUFFIXES

_PEER_NPTS = re.compile(r"NPTS\s*[=:]\s*(\d+)", re.IGNORECASE)
_PEER_DT = re.compile(r"\bDT\s*[=:]\s*(\d*\.?\d+(?:[eE][+-]?\d+)?)", re.IGNORECASE)
_PEER_DT_FALLBACK = re.compile(r"(\d*\.?\d+(?:[eE][+-]?\d+)?)\s*SEC", re.IGNORECASE)
_PEER_NPTS_DT_TRAILING = re.compile(
    r"(\d+)\s+(\d*\.?\d+(?:[eE][+-]?\d+)?)\s+NPTS\s*,?\s*DT",
    re.IGNORECASE,
)
_PEER_UNIT_TOKENS = (
    ("m/s2", ("m/sec2", "m/s2", "m/sec^2", "m/s^2")),
    ("cm/s2", ("cm/sec2", "cm/s2", "cm/sec^2", "cm/s^2", "gal")),
    ("g", ("g",)),
)


@dataclass(frozen=True)
class LoadTable:
    format_name: str
    columns: tuple[str, ...]
    rows: tuple[dict[str, str], ...]
    delimiter: str | None
    has_header: bool
    encoding: str | None
    self_describing: bool
    metadata: dict[str, Any] = field(default_factory=dict)
    warnings: tuple[str, ...] = ()


def number_or_none(value: str) -> float | None:
    text = value.strip()
    if not text:
        return None
    try:
        return float(text.replace("D", "E").replace("d", "e"))
    except ValueError:
        return None


def read_load_table(path: Path, content: bytes) -> LoadTable:
    if not content:
        raise FemCoreError("EMPTY_LOAD_FILE", "The load file is empty")
    if len(content) > MAX_LOAD_BYTES:
        raise FemCoreError("LOAD_FILE_TOO_LARGE", "The load file exceeds the 20 MiB inspection limit")

    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_LOAD_SUFFIXES:
        raise FemCoreError(
            "UNSUPPORTED_LOAD_FORMAT",
            "Supported load formats are CSV, TXT, DAT, XLSX and PEER NGA AT1/AT2",
            details={"suffix": suffix, "supported": sorted(SUPPORTED_LOAD_SUFFIXES)},
        )
    if suffix in PEER_SUFFIXES:
        return _read_peer(content)
    if suffix == ".xlsx":
        return _read_xlsx(content)
    return _read_text_table(suffix, content)


def _read_text_table(suffix: str, content: bytes) -> LoadTable:
    text, encoding = decode_engineering_text(content)
    data_lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip() and not line.lstrip().startswith(("#", "!"))
    ]
    if not data_lines:
        raise FemCoreError("NO_LOAD_DATA", "The load file contains no readable data rows")

    delimiter = _infer_delimiter(data_lines[0])
    first = _split(data_lines[0], delimiter)
    has_header = any(number_or_none(value) is None for value in first)
    columns = _unique_columns(first) if has_header else [f"column_{index}" for index in range(1, len(first) + 1)]
    raw_rows = data_lines[1:] if has_header else data_lines
    if len(raw_rows) > MAX_ROWS:
        raise FemCoreError("TOO_MANY_LOAD_ROWS", f"Load inspection is limited to {MAX_ROWS} rows")

    parsed_rows = [_split(line, delimiter) for line in raw_rows]
    inconsistent_rows = sum(len(row) != len(columns) for row in parsed_rows)
    normalized = [(row + [""] * len(columns))[: len(columns)] for row in parsed_rows]
    rows = tuple(dict(zip(columns, row, strict=True)) for row in normalized)
    warnings = ("INCONSISTENT_COLUMN_COUNT",) if inconsistent_rows else ()
    return LoadTable(
        format_name=suffix.removeprefix(".").upper(),
        columns=tuple(columns),
        rows=rows,
        delimiter=delimiter,
        has_header=has_header,
        encoding=encoding,
        self_describing=False,
        metadata={"inconsistentRowCount": inconsistent_rows},
        warnings=warnings,
    )


def _read_xlsx(content: bytes) -> LoadTable:
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise FemCoreError("INVALID_XLSX", "The XLSX workbook could not be read") from exc
    try:
        sheet = workbook[workbook.sheetnames[0]]
        raw_rows: list[list[str]] = []
        for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if index > MAX_ROWS + 1:
                raise FemCoreError("TOO_MANY_LOAD_ROWS", f"Load inspection is limited to {MAX_ROWS} rows")
            values = ["" if value is None else str(value).strip() for value in row]
            while values and not values[-1]:
                values.pop()
            if values and any(values):
                raw_rows.append(values)
    finally:
        workbook.close()

    if not raw_rows:
        raise FemCoreError("NO_LOAD_DATA", "The XLSX workbook contains no readable data rows")
    first = raw_rows[0]
    has_header = any(number_or_none(value) is None for value in first)
    columns = _unique_columns(first) if has_header else [f"column_{index}" for index in range(1, len(first) + 1)]
    body = raw_rows[1:] if has_header else raw_rows
    if len(body) > MAX_ROWS:
        raise FemCoreError("TOO_MANY_LOAD_ROWS", f"Load inspection is limited to {MAX_ROWS} rows")
    inconsistent_rows = sum(len(row) != len(columns) for row in body)
    normalized = [(row + [""] * len(columns))[: len(columns)] for row in body]
    rows = tuple(dict(zip(columns, row, strict=True)) for row in normalized)
    warnings = ("INCONSISTENT_COLUMN_COUNT",) if inconsistent_rows else ()
    return LoadTable(
        format_name="XLSX",
        columns=tuple(columns),
        rows=rows,
        delimiter=None,
        has_header=has_header,
        encoding=None,
        self_describing=False,
        metadata={"worksheet": sheet.title, "inconsistentRowCount": inconsistent_rows},
        warnings=warnings,
    )


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_XLSX_MEMBERS:
                raise FemCoreError("XLSX_TOO_MANY_MEMBERS", "The XLSX archive contains too many members")
            extracted = sum(member.file_size for member in members)
            if extracted > MAX_XLSX_EXTRACTED_BYTES:
                raise FemCoreError(
                    "XLSX_EXTRACTED_TOO_LARGE",
                    "The XLSX archive exceeds the 100 MiB extracted-size limit",
                )
    except BadZipFile as exc:
        raise FemCoreError("INVALID_XLSX", "The XLSX file is not a valid ZIP workbook") from exc


def _read_peer(content: bytes) -> LoadTable:
    text, encoding = decode_engineering_text(content)
    lines = text.splitlines()
    if not lines:
        raise FemCoreError("EMPTY_LOAD_FILE", "The PEER record is empty")

    header_lines: list[str] = []
    values: list[float] = []
    data_started = False
    for line in lines:
        stripped = line.strip()
        if not data_started:
            if not stripped:
                header_lines.append(line)
                continue
            numbers = _all_numbers_or_none(stripped)
            if numbers is None:
                header_lines.append(line)
                continue
            data_started = True
            values.extend(numbers)
            continue
        if not stripped:
            continue
        numbers = _all_numbers_or_none(stripped)
        if numbers is None:
            raise FemCoreError(
                "PEER_UNEXPECTED_TRAILING_TEXT",
                "The PEER numeric block contains unexpected trailing text",
                details={"text": stripped[:80]},
            )
        values.extend(numbers)

    if not values:
        raise FemCoreError("PEER_NO_DATA", "The PEER record contains no numeric samples")
    if len(values) > MAX_ROWS:
        raise FemCoreError("TOO_MANY_LOAD_ROWS", f"Load inspection is limited to {MAX_ROWS} rows")

    header_text = "\n".join(header_lines)
    dt = _peer_dt(header_text)
    npts = _peer_npts(header_text)
    unit = _peer_unit(header_text)
    if npts is not None and npts != len(values):
        raise FemCoreError(
            "PEER_NPTS_MISMATCH",
            "PEER declared NPTS does not match the parsed sample count",
            details={"declaredNpts": npts, "parsedSamples": len(values)},
        )

    rows = tuple(
        {"time_s": _format_number(index * dt), "acceleration": _format_number(value)}
        for index, value in enumerate(values)
    )
    descriptions = [line.strip() for line in header_lines if line.strip()]
    return LoadTable(
        format_name="PEER_NGA",
        columns=("time_s", "acceleration"),
        rows=rows,
        delimiter=None,
        has_header=True,
        encoding=encoding,
        self_describing=True,
        metadata={
            "declaredNpts": npts,
            "declaredDt": dt,
            "declaredUnit": unit,
            "declaredQuantity": "ACCELERATION",
            "description": descriptions[1] if len(descriptions) >= 2 else "",
        },
    )


def _peer_dt(header_text: str) -> float:
    match = _PEER_DT.search(header_text) or _PEER_DT_FALLBACK.search(header_text)
    if match is not None:
        dt = float(match.group(1))
    else:
        trailing = _PEER_NPTS_DT_TRAILING.search(header_text)
        if trailing is None:
            raise FemCoreError("PEER_MISSING_DT", "The PEER header does not declare DT")
        dt = float(trailing.group(2))
    if dt <= 0:
        raise FemCoreError("PEER_INVALID_DT", "The PEER DT must be positive", details={"dt": dt})
    return dt


def _peer_npts(header_text: str) -> int | None:
    match = _PEER_NPTS.search(header_text)
    if match is not None:
        return int(match.group(1))
    trailing = _PEER_NPTS_DT_TRAILING.search(header_text)
    return int(trailing.group(1)) if trailing else None


def _peer_unit(header_text: str) -> str:
    lowered = header_text.lower()
    units_index = lowered.rfind("units of")
    scope = lowered[units_index + len("units of") :] if units_index >= 0 else lowered
    for canonical, tokens in _PEER_UNIT_TOKENS:
        for token in tokens:
            if re.search(rf"(?<![a-z0-9/^]){re.escape(token)}(?![a-z0-9/^])", scope):
                return canonical
    raise FemCoreError(
        "PEER_MISSING_UNIT",
        "The PEER header does not declare a supported acceleration unit",
    )


def _all_numbers_or_none(line: str) -> list[float] | None:
    cells = [cell for cell in re.split(r"[,\s]+", line.strip()) if cell]
    if not cells:
        return None
    numbers: list[float] = []
    for cell in cells:
        value = number_or_none(cell)
        if value is None:
            return None
        numbers.append(value)
    return numbers


def _infer_delimiter(line: str) -> str:
    if "," in line:
        return ","
    if "\t" in line:
        return "\t"
    if ";" in line:
        return ";"
    return "WHITESPACE"


def _split(line: str, delimiter: str) -> list[str]:
    if delimiter == "WHITESPACE":
        return re.split(r"\s+", line.strip())
    return [part.strip() for part in line.split(delimiter)]


def _unique_columns(values: list[str]) -> list[str]:
    result: list[str] = []
    counts: dict[str, int] = {}
    for index, raw in enumerate(values, start=1):
        base = raw.strip() or f"column_{index}"
        counts[base] = counts.get(base, 0) + 1
        result.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return result


def _format_number(value: float) -> str:
    return format(value, ".15g")
